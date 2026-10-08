"""Inventario de encabezados de los archivos internos (apoyo a los pasos 3.1.2 y 3.3.2).

Lista, por archivo y por hoja, SOLO los nombres de las columnas. Es el paso
previo a declarar los esquemas reales en `src/common/schemas.py`.

Garantía frente a la regla de `CLAUDE.md` («nunca imprimir ni commitear
`data/raw/`»): de cada hoja sale como máximo una fila, la del encabezado, y
solo si esa fila *parece* un encabezado. Hay dos barreras:

  1. `elegir_encabezado` puntúa las primeras filas por tipo de celda, no por
     cuántas están llenas: un encabezado es texto y la fila de abajo no lo es.
     Si ninguna fila pasa el umbral, la hoja se marca `encabezado_dudoso` y se
     publican el número de columnas y sus letras, pero NINGÚN nombre.
  2. `_valor_encabezado` descarta celda por celda lo que parece dato aunque la
     fila haya pasado: fechas, decimales y números largos salen como `omitido`.
     Se conservan los números cortos (`2020`, `15`), que sí son encabezados
     legítimos en los cuadros por año.

Del resto de las filas solo se evalúa el tipo de cada celda para puntuar; el
valor se descarta en el mismo paso y nunca llega a la salida.

Además de listar, resuelve lo que rompe a los lectores ingenuos:
  * hojas múltiples y filas de título antes del encabezado;
  * codificación y delimitador de los CSV (utf-8 con BOM, cp1252, `;`, `|`);
  * `.xls` antiguo (vía xlrd) y archivos con la extensión equivocada;
  * encabezados repetidos y columnas sin nombre;
  * contraste contra `schemas.FUENTES` cuando la fuente ya está declarada.

Uso:
  python -m src.ingest.encabezados --origen ../DATASETS
  python -m src.ingest.encabezados --origen data/raw --sin-nombres

`--sin-nombres` publica solo el recuento y las letras de columna, sin un solo
nombre: es la opción para cuando la salida va a un repositorio público.

La salida (JSON + Markdown) solo contiene metadatos. Los archivos de datos no.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import unicodedata
from pathlib import Path

from src.common import config, schemas

EXTENSIONES_TABLA = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}
CODIFICACIONES = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
DELIMITADORES = (",", ";", "|", "\t")
FILAS_INSPECCION = 20          # filas iniciales donde se busca el encabezado
MUESTRA_BYTES = 64 * 1024      # para adivinar codificación y delimitador
MAX_MB_CONTEO = 500            # sobre este tamaño no se cuentan filas de CSV
FILAS_MAX_EXCEL = 1_048_576    # tope de una hoja: max_row cercano a esto no es fiable

UMBRAL_TEXTO = 0.5             # fracción mínima de celdas de texto en un encabezado
UMBRAL_CONTRASTE = 0.25        # cuánto más texto que la fila de abajo
FRACCION_ANCHO = 0.6           # ancho mínimo frente a la fila más llena
LARGO_NUMERO_OK = 6            # «2020» pasa; «896902.36» no

SALIDA_JSON = "docs/evidencias/encabezados_raw.json"
SALIDA_MD = "docs/evidencias/encabezados_raw.md"

_RE_CODIGO = re.compile(r"\b([A-Za-z]{3,4})[-_ ]?(\d{2})\b")
_RE_FECHA_TEXTO = re.compile(r"^\d{1,4}[-/]\d{1,2}[-/]\d{1,4}")
_RE_DOC_LARGO = re.compile(r"\d{8,}")                    # RUC (11), DNI (8), documento
_RE_RAZON_SOCIAL = re.compile(
    r"(\bS\.A\.C?\.?|\bE\.I\.R\.L\.?|\bS\.R\.L\.?|\sS\.?A\.?C?\.?\s*$)", re.IGNORECASE
)
_RE_PLACA = re.compile(r"^[A-Za-z0-9]{2,4}[-\s][A-Za-z0-9]{2,4}$")
_LIMPIEZA_NUMERO = str.maketrans({c: "" for c in "  S/$%()',"})


# --- tipos de celda -----------------------------------------------------------
def tipo_celda(valor: object) -> str:
    """Clasifica una celda en vacia / texto / numero / fecha. No devuelve el valor."""
    if valor is None:
        return "vacia"
    if isinstance(valor, bool):
        return "texto"
    if isinstance(valor, (dt.datetime, dt.date, dt.time)):
        return "fecha"
    if isinstance(valor, (int, float)):
        return "numero"
    s = str(valor).strip()
    if not s:
        return "vacia"
    try:
        float(s.translate(_LIMPIEZA_NUMERO))
        return "numero"
    except ValueError:
        pass
    if _RE_FECHA_TEXTO.match(s):
        return "fecha"
    return "texto"


def _no_vacia(valor: object) -> bool:
    return tipo_celda(valor) != "vacia"


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    return " ".join(str(valor).split())


def normalizar_nombre(nombre: str) -> str:
    """Clave de comparación: sin tildes, sin espacios dobles, en mayúsculas."""
    sin_tilde = "".join(
        c for c in unicodedata.normalize("NFKD", nombre) if not unicodedata.combining(c)
    )
    return " ".join(sin_tilde.split()).upper()


# --- eleccion de la fila de encabezado ----------------------------------------
def _perfil(fila: list) -> tuple[int, float]:
    """(celdas llenas, fracción de esas celdas que es texto)."""
    tipos = [tipo_celda(v) for v in fila]
    llenas = sum(1 for t in tipos if t != "vacia")
    if not llenas:
        return 0, 0.0
    return llenas, sum(1 for t in tipos if t == "texto") / llenas


def elegir_encabezado(filas: list[list]) -> tuple[int, list, bool]:
    """Devuelve (numero de fila 1-based, valores de esa fila, encabezado_dudoso).

    Puntúa por tipo, no por cantidad: una fila de datos numéricos puede tener
    más celdas llenas que el encabezado y no por eso es el encabezado.
    """
    if not filas:
        return 1, [], True
    perfiles = [_perfil(f) for f in filas]
    ancho_max = max(p[0] for p in perfiles)
    if ancho_max == 0:
        return 1, [], True

    minimo = max(2, int(ancho_max * FRACCION_ANCHO))
    mejor_clave, mejor_i = None, None
    for i, (llenas, ratio) in enumerate(perfiles):
        if llenas < minimo:
            continue
        hay_siguiente = i + 1 < len(perfiles)
        contrasta = (
            hay_siguiente and (ratio - perfiles[i + 1][1]) >= UMBRAL_CONTRASTE
        )
        clave = (
            ratio >= UMBRAL_TEXTO,                       # parece encabezado
            contrasta,                                   # y la fila de abajo no
            llenas,
            -i,                                          # a igualdad, la primera
        )
        if mejor_clave is None or clave > mejor_clave:
            mejor_clave, mejor_i = clave, i

    if mejor_i is None:                                  # ninguna alcanza el ancho
        mejor_i = max(range(len(perfiles)), key=lambda i: (perfiles[i][0], -i))
    return mejor_i + 1, list(filas[mejor_i]), perfiles[mejor_i][1] < UMBRAL_TEXTO


def parece_identificador(texto: str) -> bool:
    """¿La celda es un documento, una razón social o una placa?

    Un encabezado no contiene el RUC de nadie. Si aparece uno, la fila es un
    registro de datos y la hoja entera se trata como sin encabezado, porque
    `CLAUDE.md` prohíbe que un RUC o un DNI salga en cualquier salida.
    """
    t = texto.strip()
    if not t:
        return False
    if _RE_DOC_LARGO.search(t) or _RE_RAZON_SOCIAL.search(t):
        return True
    if len(t) <= 9 and _RE_PLACA.fullmatch(t):
        return any(c.isdigit() for c in t) and any(c.isalpha() for c in t)
    return False


def _valor_encabezado(valor: object) -> tuple[str, bool]:
    """(nombre publicable, omitido). Descarta la celda si parece un dato."""
    t = tipo_celda(valor)
    if t == "vacia":
        return "", False
    if t == "texto":
        return _texto(valor), False
    if t == "numero":
        s = _texto(valor)
        if len(s) <= LARGO_NUMERO_OK and "." not in s:
            return s, False                              # 2020, 15: encabezado legítimo
    return "", True                                      # fecha, decimal o número largo


def _letra_columna(pos: int) -> str:
    letras = ""
    while pos > 0:
        pos, resto = divmod(pos - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def _columnas(valores: list, dudoso: bool, sin_nombres: bool) -> list[dict]:
    cols = []
    for i, v in enumerate(valores, start=1):
        if dudoso or sin_nombres:
            nombre, omitido = "", _no_vacia(v)
        else:
            nombre, omitido = _valor_encabezado(v)
        cols.append({"pos": i, "col": _letra_columna(i), "nombre": nombre, "omitido": omitido})
    while cols and not cols[-1]["nombre"] and not cols[-1]["omitido"]:
        cols.pop()                                       # cola de columnas vacías
    return cols


def _avisos(cols: list[dict]) -> dict:
    nombres = [c["nombre"] for c in cols if c["nombre"]]
    vistos, duplicadas = set(), []
    for n in nombres:
        clave = normalizar_nombre(n)
        if clave in vistos and n not in duplicadas:
            duplicadas.append(n)
        vistos.add(clave)
    return {
        "duplicadas": duplicadas,
        "sin_nombre": [c["col"] for c in cols if not c["nombre"] and not c["omitido"]],
        "omitidas": [c["col"] for c in cols if c["omitido"]],
    }


def _hoja(nombre: str | None, filas: list[list], total_filas: int | None,
          sin_nombres: bool, **extra) -> dict:
    fila_enc, valores, dudoso = elegir_encabezado(filas)
    motivo = "tipo_de_celda" if dudoso else None
    if not dudoso and any(parece_identificador(_texto(v)) for v in valores):
        dudoso, motivo = True, "identificador_en_la_fila"
    cols = _columnas(valores, dudoso, sin_nombres)
    filas_datos = None
    if total_filas is not None and total_filas < FILAS_MAX_EXCEL - 10:
        filas_datos = max(total_filas - fila_enc, 0)
    return {
        "hoja": nombre,
        "fila_encabezado": fila_enc,
        "encabezado_dudoso": dudoso,
        "motivo_dudoso": motivo,
        "n_columnas": len(cols),
        "filas_datos": filas_datos,
        "columnas": cols,
        **_avisos(cols),
        **extra,
    }


def _primeras(iterable, n: int = FILAS_INSPECCION) -> list[list]:
    filas = []
    for i, fila in enumerate(iterable):
        if i >= n:
            break
        filas.append(list(fila))
    return filas


# --- CSV ----------------------------------------------------------------------
def detectar_codificacion(ruta: Path) -> str:
    bruto = ruta.open("rb").read(MUESTRA_BYTES)
    for cod in CODIFICACIONES:
        try:
            bruto.decode(cod)
            return cod
        except UnicodeDecodeError:
            continue
    return "latin-1"


def detectar_delimitador(ruta: Path, codificacion: str) -> str:
    """Elige el delimitador que se repite el mismo número de veces en más líneas.

    Mirar solo la primera línea falla cuando el archivo empieza con un título
    sin delimitadores, que es el caso de varias entregas reales.
    """
    with ruta.open("r", encoding=codificacion, errors="replace") as f:
        muestra = f.read(MUESTRA_BYTES)
    lineas = [l for l in muestra.splitlines() if l.strip()][:FILAS_INSPECCION]

    mejor, mejor_puntaje = None, (0, 0)
    for d in DELIMITADORES:
        cuentas = [l.count(d) for l in lineas if l.count(d) > 0]
        if not cuentas:
            continue
        modo = max(set(cuentas), key=cuentas.count)
        puntaje = (cuentas.count(modo), modo)            # líneas consistentes, luego columnas
        if puntaje > mejor_puntaje:
            mejor, mejor_puntaje = d, puntaje
    if mejor:
        return mejor
    try:
        return csv.Sniffer().sniff(muestra, delimiters="".join(DELIMITADORES)).delimiter
    except csv.Error:
        return ","


def _contar_filas_csv(ruta: Path, codificacion: str) -> int | None:
    if ruta.stat().st_size > MAX_MB_CONTEO * 1024 * 1024:
        return None
    with ruta.open("r", encoding=codificacion, errors="replace", newline="") as f:
        return sum(1 for linea in f if linea.strip())


def leer_csv(ruta: Path, sin_nombres: bool = False) -> list[dict]:
    cod = detectar_codificacion(ruta)
    delim = detectar_delimitador(ruta, cod)
    with ruta.open("r", encoding=cod, errors="replace", newline="") as f:
        filas = _primeras(csv.reader(f, delimiter=delim))
    return [_hoja(None, filas, _contar_filas_csv(ruta, cod), sin_nombres,
                  codificacion=cod, delimitador="\\t" if delim == "\t" else delim)]


# --- Excel --------------------------------------------------------------------
def _hojas_de_libro(fuente, sin_nombres: bool) -> list[dict]:
    """`fuente` es una ruta o un archivo binario abierto.

    openpyxl rechaza un libro por su extensión, no por su contenido, así que un
    .xlsx mal nombrado solo se abre pasándole el descriptor ya abierto.
    """
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(fuente, read_only=True, data_only=True)
    try:
        hojas = []
        for ws in wb.worksheets:
            filas = _primeras(
                ws.iter_rows(min_row=1, max_row=FILAS_INSPECCION, values_only=True)
            )
            hojas.append(_hoja(ws.title, filas, ws.max_row, sin_nombres,
                               codificacion=None, delimitador=None))
        return hojas
    finally:
        wb.close()


def leer_xlsx(ruta: Path, sin_nombres: bool = False) -> list[dict]:
    return _hojas_de_libro(ruta, sin_nombres)


def leer_xls(ruta: Path, sin_nombres: bool = False) -> list[dict]:
    """`.xls` BIFF con xlrd, y rescate de los archivos con extensión equivocada."""
    cabeza = ruta.open("rb").read(16)
    if cabeza[:2] == b"PK":                              # es un .xlsx mal nombrado
        with ruta.open("rb") as fh:
            return _hojas_de_libro(fh, sin_nombres)
    if cabeza.lstrip()[:1] == b"<":                      # HTML/XML de exportación
        raise RuntimeError(
            "es HTML con extensión .xls (exportación del ERP), no un Excel binario. "
            "Ábrelo en Excel y guárdalo como .xlsx."
        )
    try:
        import xlrd  # noqa: PLC0415
    except ImportError:
        raise RuntimeError(
            "formato .xls antiguo: falta xlrd. Instálalo con "
            "`pip install xlrd` (ya está en requirements.txt)."
        ) from None

    wb = xlrd.open_workbook(ruta)
    try:
        hojas = []
        for nombre in wb.sheet_names():
            ws = wb.sheet_by_name(nombre)
            filas = [ws.row_values(i) for i in range(min(ws.nrows, FILAS_INSPECCION))]
            hojas.append(_hoja(nombre, filas, ws.nrows, sin_nombres,
                               codificacion=None, delimitador=None))
        return hojas
    finally:
        wb.release_resources()


LECTORES = {
    ".xlsx": leer_xlsx, ".xlsm": leer_xlsx, ".xls": leer_xls,
    ".csv": leer_csv, ".tsv": leer_csv, ".txt": leer_csv,
}


# --- codigo de fuente ---------------------------------------------------------
def inferir_codigo(ruta_relativa: Path) -> str | None:
    """Deduce el código de fuente (`aba_01`) de la ruta o del nombre del archivo.

    Acepta `abastecimiento/aba_01/x.xlsx`, `ABA-01 ordenes.xlsx` y `aba01.csv`.
    Devuelve None cuando el archivo no sigue ninguna nomenclatura de código.
    """
    for parte in (*reversed(ruta_relativa.parts[:-1]), ruta_relativa.stem):
        m = _RE_CODIGO.search(parte)
        if m:
            return f"{m.group(1).lower()}_{m.group(2)}"
    return None


def inferir_proceso(ruta_relativa: Path) -> str | None:
    partes = ruta_relativa.parts[:-1]
    return partes[0] if partes else None


# --- contraste con los esquemas declarados ------------------------------------
def _contrastar_hoja(declaradas: list[str], hoja: dict) -> dict:
    halladas = [c["nombre"] for c in hoja["columnas"] if c["nombre"]]
    claves_halladas = {normalizar_nombre(n) for n in halladas}
    claves_declaradas = {normalizar_nombre(n) for n in declaradas}
    faltan = [n for n in declaradas if normalizar_nombre(n) not in claves_halladas]
    return {
        "hoja": hoja["hoja"],
        "coincidencias": len(declaradas) - len(faltan),
        "declaradas_no_encontradas": faltan,
        "encontradas_no_declaradas": [
            n for n in dict.fromkeys(halladas) if normalizar_nombre(n) not in claves_declaradas
        ],
    }


def contrastar(codigo: str | None, hojas: list[dict]) -> dict | None:
    """Compara los encabezados hallados con `schemas.FUENTES[codigo]`.

    El contraste es por hoja y se reporta el de la hoja que más columnas
    declaradas reconoce: en las entregas reales la hoja de datos convive con
    hojas de totales o de notas, y mezclarlas ensucia el resultado.
    """
    if not codigo or codigo not in schemas.FUENTES:
        return None
    declaradas = list(schemas.FUENTES[codigo].columnas)
    if not hojas:
        return {"hoja": None, "coincidencias": 0,
                "declaradas_no_encontradas": declaradas, "encontradas_no_declaradas": []}
    return max((_contrastar_hoja(declaradas, h) for h in hojas),
               key=lambda c: c["coincidencias"])


# --- recorrido ----------------------------------------------------------------
def inventariar(origen: Path, sin_nombres: bool = False) -> dict:
    origen = Path(origen)
    if not origen.exists():
        raise FileNotFoundError(
            f"No existe la carpeta origen '{origen}'. "
            "Comprueba la ruta o usa --origen data/sample/entrega."
        )
    if not origen.is_dir():
        raise NotADirectoryError(f"La ruta de origen '{origen}' no es un directorio válido.")

    archivos = []
    for f in sorted(origen.rglob("*")):
        if not f.is_file() or f.name.startswith(("~$", "_", ".")):
            continue
        ext = f.suffix.lower()
        if ext not in EXTENSIONES_TABLA:
            continue
        rel = f.relative_to(origen)
        ficha: dict = {
            "ruta": rel.as_posix(),
            "proceso": inferir_proceso(rel),
            "codigo": inferir_codigo(rel),
            "extension": ext,
            "bytes": f.stat().st_size,
            "legible": True,
            "motivo": None,
            "hojas": [],
        }
        try:
            ficha["hojas"] = LECTORES[ext](f, sin_nombres)
        except Exception as e:                           # ningún archivo detiene el inventario
            ficha["legible"] = False
            ficha["motivo"] = f"{type(e).__name__}: {e}"
        ficha["contraste_esquema"] = (
            None if sin_nombres else contrastar(ficha["codigo"], ficha["hojas"])
        )
        archivos.append(ficha)

    revisar = [
        {"ruta": a["ruta"], "hoja": h["hoja"], "n_columnas": h["n_columnas"],
         "motivo": h["motivo_dudoso"]}
        for a in archivos for h in a["hojas"] if h["encabezado_dudoso"]
    ]
    con_esquema = sum(1 for a in archivos if a["contraste_esquema"] is not None)
    return {
        "generado_utc": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "origen": Path(origen).as_posix(),
        "sin_nombres": sin_nombres,
        "resumen": {
            "archivos": len(archivos),
            "hojas": sum(len(a["hojas"]) for a in archivos),
            "ilegibles": sum(1 for a in archivos if not a["legible"]),
            "hojas_sin_encabezado_claro": len(revisar),
            "con_esquema_declarado": con_esquema,
            "sin_esquema_declarado": len(archivos) - con_esquema,
            "fuentes_declaradas_en_schemas": len(schemas.FUENTES),
        },
        "hojas_para_revisar": revisar,
        "archivos": archivos,
    }


# --- salidas ------------------------------------------------------------------
def a_markdown(inv: dict) -> str:
    r = inv["resumen"]
    lineas = [
        "# Encabezados reales de las fuentes internas",
        "",
        f"Generado: {inv['generado_utc']} · origen: `{inv['origen']}`",
        "",
        (f"- Archivos: **{r['archivos']}** ({r['ilegibles']} ilegibles) · "
         f"hojas: **{r['hojas']}**"),
        f"- Hojas sin encabezado claro (sin nombres publicados): **{r['hojas_sin_encabezado_claro']}**",
        (f"- Con esquema declarado en `schemas.FUENTES`: **{r['con_esquema_declarado']}** · "
         f"sin declarar: **{r['sin_esquema_declarado']}**"),
        "",
        ("Solo nombres de columna. Las celdas que parecen dato (fechas, decimales, "
         "números largos) salen como `omitida`, y de una hoja sin encabezado claro "
         "no se publica ningún nombre."),
        "",
        "## Resumen",
        "",
        "| Archivo | Código | Hoja | Fila enc. | Cols. | Filas | Avisos |",
        "|---|---|---|---|---|---|---|",
    ]
    for a in inv["archivos"]:
        if not a["legible"]:
            lineas.append(
                f"| `{a['ruta']}` | {a['codigo'] or '—'} | — | — | — | — | "
                f"**ilegible**: {a['motivo']} |"
            )
            continue
        for h in a["hojas"]:
            avisos = []
            if h["encabezado_dudoso"]:
                avisos.append("**sin encabezado claro**")
            if h["duplicadas"]:
                avisos.append(f"duplicadas: {', '.join(h['duplicadas'])}")
            if h["sin_nombre"]:
                avisos.append(f"sin nombre: {', '.join(h['sin_nombre'])}")
            if h["omitidas"]:
                avisos.append(f"omitidas por parecer dato: {', '.join(h['omitidas'])}")
            if h["fila_encabezado"] > 1:
                avisos.append(f"{h['fila_encabezado'] - 1} fila(s) de título")
            if h["delimitador"] and h["delimitador"] != ",":
                avisos.append(f"delimitador `{h['delimitador']}`")
            lineas.append(
                f"| `{a['ruta']}` | {a['codigo'] or '—'} | {h['hoja'] or '—'} | "
                f"{h['fila_encabezado']} | {h['n_columnas']} | "
                f"{'?' if h['filas_datos'] is None else h['filas_datos']} | "
                f"{'; '.join(avisos) or '—'} |"
            )

    if inv["hojas_para_revisar"]:
        lineas += [
            "",
            "## Hojas que necesitan una mirada humana",
            "",
            ("Ninguna fila de estas hojas parece un encabezado, así que no se publica "
             "ningún nombre. Hay que abrirlas y decir en qué fila empieza la tabla "
             "(o descartarlas si son cuadros de totales)."),
            "",
            "| Archivo | Hoja | Cols. | Por qué |",
            "|---|---|---|---|",
        ]
        motivos = {
            "tipo_de_celda": "ninguna fila parece encabezado",
            "identificador_en_la_fila": "la fila traía un documento, razón social o placa",
        }
        for h in inv["hojas_para_revisar"]:
            lineas.append(
                f"| `{h['ruta']}` | {h['hoja'] or '—'} | {h['n_columnas']} | "
                f"{motivos.get(h['motivo'], h['motivo'] or '—')} |"
            )

    lineas += ["", "## Columnas por archivo", ""]
    for a in inv["archivos"]:
        lineas += [f"### `{a['ruta']}`", ""]
        if not a["legible"]:
            lineas += [f"No se pudo leer: {a['motivo']}", ""]
            continue
        for h in a["hojas"]:
            if h["hoja"]:
                cab = f"**Hoja `{h['hoja']}`** (encabezado en la fila {h['fila_encabezado']})"
            else:
                cab = (f"Codificación `{h['codificacion']}`, delimitador "
                       f"`{h['delimitador']}`, encabezado en la fila {h['fila_encabezado']}")
            lineas += [cab, ""]
            if h["encabezado_dudoso"]:
                lineas += [
                    (f"Sin encabezado claro: {h['n_columnas']} columnas "
                     f"(`A`-`{_letra_columna(h['n_columnas'])}`), nombres no publicados."),
                    "",
                ]
                continue
            lineas += ["| # | Col | Nombre de la columna |", "|---|---|---|"]
            for c in h["columnas"]:
                if c["omitido"]:
                    nombre = "*(omitida: parece dato)*"
                else:
                    nombre = c["nombre"] or "*(vacía)*"
                lineas.append(f"| {c['pos']} | {c['col']} | {nombre} |")
            lineas.append("")
        ce = a["contraste_esquema"]
        if ce:
            hoja_txt = f" · hoja `{ce['hoja']}`" if ce["hoja"] else ""
            lineas += [f"Contraste con `schemas.FUENTES['{a['codigo']}']`{hoja_txt}:", ""]
            lineas.append(
                "- Declaradas que no aparecen: "
                + (", ".join(f"`{n}`" for n in ce["declaradas_no_encontradas"]) or "ninguna")
            )
            lineas.append(
                "- En el archivo y sin declarar: "
                + (", ".join(f"`{n}`" for n in ce["encontradas_no_declaradas"]) or "ninguna")
            )
            lineas.append("")
    return "\n".join(lineas) + "\n"


def _resolver(destino: str | Path) -> Path:
    p = Path(destino)
    return p if p.is_absolute() else config.RAIZ_REPO / p


def guardar(inv: dict, salida_json: str | Path, salida_md: str | Path) -> tuple[Path, Path]:
    pj, pm = _resolver(salida_json), _resolver(salida_md)
    for p in (pj, pm):
        p.parent.mkdir(parents=True, exist_ok=True)
    pj.write_text(json.dumps(inv, ensure_ascii=False, indent=2), encoding="utf-8")
    pm.write_text(a_markdown(inv), encoding="utf-8")
    return pj, pm


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Lista archivo, hoja y columnas de una carpeta de datos, sin mostrar datos."
    )
    ap.add_argument("--origen", type=Path, default=Path("data/raw"))
    ap.add_argument("--salida", default=SALIDA_JSON)
    ap.add_argument("--salida-md", default=SALIDA_MD)
    ap.add_argument("--sin-nombres", action="store_true",
                    help="Publicar solo recuentos y letras de columna, ningún nombre.")
    ap.add_argument("--imprimir", action="store_true",
                    help="Volcar el Markdown completo a la consola.")
    a = ap.parse_args(argv)

    inv = inventariar(a.origen, a.sin_nombres)
    pj, pm = guardar(inv, a.salida, a.salida_md)
    r = inv["resumen"]

    if a.imprimir:
        print(a_markdown(inv))
        return 0

    print(f"Archivos: {r['archivos']} ({r['ilegibles']} ilegibles) · hojas: {r['hojas']}")
    print(f"Hojas sin encabezado claro: {r['hojas_sin_encabezado_claro']}")
    print(f"Con esquema declarado: {r['con_esquema_declarado']} · "
          f"sin declarar: {r['sin_esquema_declarado']}")
    for arch in inv["archivos"]:
        if not arch["legible"]:
            print(f"  ILEGIBLE  {arch['ruta']} -> {arch['motivo']}")
    print(f"Salida: {pj}")
    print(f"        {pm}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
