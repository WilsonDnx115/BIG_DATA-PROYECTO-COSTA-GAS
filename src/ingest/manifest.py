"""Manifiesto de carga (paso 3.2.1): base del linaje Bronze → Gold (5.2).

Cada archivo subido a Bronze queda registrado con fuente, ruta, tamaño,
SHA-256, fecha de carga y número de filas (si es tabular).
"""
from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path

from src.common import lake_io, rutas


def sha256(ruta: Path, bloque: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for trozo in iter(lambda: f.read(bloque), b""):
            h.update(trozo)
    return h.hexdigest()


def contar_filas(ruta: Path) -> int | None:
    suf = ruta.suffix.lower()
    if suf == ".csv":
        with open(ruta, "rb") as f:
            return max(sum(1 for _ in f) - 1, 0)
    if suf in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook  # noqa: PLC0415

            wb = load_workbook(ruta, read_only=True)
            return sum(max(ws.max_row - 1, 0) for ws in wb.worksheets)
        except Exception:
            return None
    return None


def entrada(local: Path, fuente: str, destino: str, fecha_carga: str) -> dict:
    return {
        "fuente": fuente,
        "archivo_original": local.name,
        "ruta_bronze": destino,
        "bytes": local.stat().st_size,
        "sha256": sha256(local),
        "filas": contar_filas(local),
        "fecha_carga": fecha_carga,
        "registrado_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


class Manifiesto:
    def __init__(self, fecha_carga: str | None = None):
        self.fecha_carga = fecha_carga or date.today().isoformat()
        self.entradas: list[dict] = []

    def agregar(self, local: Path, fuente: str, destino: str) -> dict:
        e = entrada(local, fuente, destino, self.fecha_carga)
        self.entradas.append(e)
        return e

    def guardar(self) -> str:
        destino = rutas.ruta("_manifests", f"fecha_carga={self.fecha_carga}", "manifest.json")
        previo = []
        if not rutas.es_s3() and Path(destino).exists():
            previo = json.loads(Path(destino).read_text(encoding="utf-8"))
        vistos = {(e["ruta_bronze"], e["sha256"]) for e in previo}
        todo = previo + [e for e in self.entradas if (e["ruta_bronze"], e["sha256"]) not in vistos]
        lake_io.escribir_json(todo, destino)
        return destino
