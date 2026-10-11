"""Lectura de hojas de Excel con esquema declarado (hallazgos C5 y C6).

* No se infiere nada: la fila del encabezado y las columnas vienen del contrato.
* Corte de hojas infladas: muchas hojas arrastran fórmulas hasta la fila 1 048 576.
  Se deja de leer tras `MAX_FILAS_VACIAS` filas vacías seguidas (S3: el costo de
  recorrer filas inútiles supera al trabajo útil). `leidas` y `efectivas` quedan
  en la tabla de control para medir el antes y el después.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from src.common import config
from src.silver.contratos import Fuente, normalizar_encabezado


class HojaRechazada(ValueError):
    """La hoja no cumple el contrato (falta una columna declarada)."""


@dataclass
class Lectura:
    hoja: str
    datos: pd.DataFrame          # columnas con nombre estándar + fila
    leidas: int                  # filas recorridas bajo el encabezado
    efectivas: int               # filas con algún dato en las columnas del contrato
    cortada: bool                # se detuvo por el corte de filas vacías


def _vacia(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and not v.strip())


def nombres_hojas(ruta: str) -> list[str]:
    if Path(ruta).suffix.lower() == ".xls":
        import xlrd  # noqa: PLC0415

        return xlrd.open_workbook(ruta, on_demand=True).sheet_names()
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _filas_xlsx(ruta: str, hoja: str, desde: int):
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        yield from wb[hoja].iter_rows(min_row=desde, values_only=True)
    finally:
        wb.close()


def _filas_xls(ruta: str, hoja: str, desde: int):
    import xlrd  # noqa: PLC0415

    libro = xlrd.open_workbook(ruta, on_demand=True)
    h = libro.sheet_by_name(hoja)
    for i in range(desde - 1, h.nrows):
        fila = []
        for celda in h.row(i):
            if celda.ctype == xlrd.XL_CELL_DATE:
                try:
                    fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode))
                except Exception:
                    fila.append(None)
            elif celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                fila.append(None)
            else:
                fila.append(celda.value)
        yield tuple(fila)


def leer_hoja(ruta: str, hoja: str, fuente: Fuente, max_vacias: int | None = None) -> Lectura:
    """Lee una hoja según el contrato y devuelve las columnas declaradas con su nombre estándar."""
    max_vacias = config.MAX_FILAS_VACIAS if max_vacias is None else max_vacias
    filas = (_filas_xls if Path(ruta).suffix.lower() == ".xls" else _filas_xlsx)(ruta, hoja, fuente.fila_encabezado)
    try:
        encabezado = next(filas)
    except StopIteration:
        raise HojaRechazada(f"{hoja}: no tiene la fila de encabezado {fuente.fila_encabezado}") from None

    posiciones: dict[str, int] = {}
    for i, celda in enumerate(encabezado):
        nombre = fuente.columnas.get(normalizar_encabezado(celda))
        if nombre and nombre not in posiciones:            # ante encabezados repetidos se toma el primero
            posiciones[nombre] = i
    faltan = sorted(set(fuente.columnas.values()) - set(posiciones))
    if faltan:
        raise HojaRechazada(f"{hoja}: no cumple el contrato {fuente.id}; faltan las columnas {faltan}")

    registros, leidas, vacias, cortada = [], 0, 0, False
    for n, fila in enumerate(filas, start=fuente.fila_encabezado + 1):
        leidas += 1
        valores = {nombre: (fila[i] if i < len(fila) else None) for nombre, i in posiciones.items()}
        if all(_vacia(v) for v in valores.values()):
            vacias += 1
            if vacias >= max_vacias:
                cortada = True
                break
            continue
        vacias = 0
        valores["fila"] = n
        registros.append(valores)

    datos = pd.DataFrame(registros, columns=[*posiciones, "fila"])
    return Lectura(hoja=hoja, datos=datos, leidas=leidas, efectivas=len(datos), cortada=cortada)


def a_texto(v) -> str | None:
    """Celda → texto limpio. Los números enteros guardados como float pierden el '.0'."""
    if _vacia(v):
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return " ".join(str(v).split())
